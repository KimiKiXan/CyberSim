"""Build CyberSim documentation in TXT, DOCX and PDF.

Usage:
    python docs/build_docs.py                      # English (DOCUMENTATION.md)
    python docs/build_docs.py --lang ru            # Russian  (DOCUMENTATION_RU.md)
    python docs/build_docs.py --lang en --lang ru  # Both
    python docs/build_docs.py --source path.md --prefix MyName

Outputs are written next to the source markdown.
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

ROOT = Path(__file__).resolve().parent

LANGS = {
    "en": {
        "source":     ROOT / "DOCUMENTATION.md",
        "prefix":     "CyberSim_Documentation",
        "title":      "CyberSim — Official Documentation",
        "subtitle":   "Autonomous LLM-Driven Cyberattack Simulation Framework — Version 0.1.0",
        "footer":     "CyberSim Documentation — v0.1.0",
        "page_word":  "Page",
    },
    "ru": {
        "source":     ROOT / "DOCUMENTATION_RU.md",
        "prefix":     "CyberSim_Documentation_RU",
        "title":      "CyberSim — Официальная документация",
        "subtitle":   "Автономный фреймворк симуляции кибератак на базе LLM — Версия 0.1.0",
        "footer":     "Документация CyberSim — v0.1.0",
        "page_word":  "Страница",
    },
}


# ---------------------------------------------------------------- markdown AST
@dataclass
class Block:
    kind: str           # h1 / h2 / h3 / h4 / p / ul / ol / code / table / hr
    content: object     # str, list[str] (lists), list[list[str]] (tables), …
    lang: str = ""


def parse_markdown(text: str) -> list[Block]:
    """Tiny markdown parser sufficient for our internal documentation."""
    blocks: list[Block] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        stripped = raw.rstrip()

        # fenced code block
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            body: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            blocks.append(Block("code", "\n".join(body), lang=lang))
            continue

        # blank line — skip, blocks already separate themselves
        if not stripped.strip():
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}\s*$", stripped):
            blocks.append(Block("hr", ""))
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            blocks.append(Block(f"h{level}", m.group(2).strip()))
            i += 1
            continue

        # table
        if "|" in stripped and i + 1 < len(lines) and re.match(r"^\s*\|?[\s\-:|]+\|[\s\-:|]+", lines[i + 1]):
            header = _split_row(stripped)
            i += 2  # skip header + divider
            rows = [header]
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            blocks.append(Block("table", rows))
            continue

        # unordered list
        if re.match(r"^\s*[-*]\s+", stripped):
            items: list[str] = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i] or ""):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).rstrip())
                i += 1
            blocks.append(Block("ul", items))
            continue

        # ordered list
        if re.match(r"^\s*\d+\.\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i] or ""):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]).rstrip())
                i += 1
            blocks.append(Block("ol", items))
            continue

        # paragraph: collect until blank line / list / heading / table / fence
        para_lines: list[str] = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip():
                break
            if re.match(r"^(#{1,6})\s+", nxt):
                break
            if nxt.startswith("```"):
                break
            if re.match(r"^\s*([-*]|\d+\.)\s+", nxt):
                break
            if "|" in nxt and i + 1 < len(lines) and re.match(r"^\s*\|?[\s\-:|]+\|[\s\-:|]+", lines[i + 1]):
                break
            para_lines.append(nxt.rstrip())
            i += 1
        blocks.append(Block("p", " ".join(_strip_inline(p) for p in para_lines)))
    return blocks


def _split_row(line: str) -> list[str]:
    parts = line.strip().strip("|").split("|")
    return [_strip_inline(p.strip()) for p in parts]


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITAL_RE = re.compile(r"(?<![\*])\*(?!\*)(.+?)\*(?!\*)")
_CODE_RE = re.compile(r"`([^`]+)`")
_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _strip_inline(text: str) -> str:
    """Return readable text — strip markdown emphasis markers for TXT/DOCX."""
    text = _LINK_RE.sub(r"\1 (\2)", text)
    text = _BOLD_RE.sub(r"\1", text)
    text = _ITAL_RE.sub(r"\1", text)
    text = _CODE_RE.sub(r"\1", text)
    return text


# ----------------------------------------------------------------------- TXT
def render_txt(blocks: list[Block]) -> str:
    out: list[str] = []
    for b in blocks:
        if b.kind == "hr":
            out += ["", "-" * 78, ""]
            continue
        if b.kind.startswith("h") and b.kind[1:].isdigit():
            level = int(b.kind[1:])
            text = b.content  # type: ignore[assignment]
            if level == 1:
                bar = "=" * max(60, len(text))
                out += [bar, str(text).upper(), bar, ""]
            elif level == 2:
                bar = "-" * max(50, len(text))
                out += ["", str(text), bar, ""]
            else:
                out += ["", "## " + str(text), ""]
        elif b.kind == "p":
            out += [_wrap(str(b.content), 90), ""]
        elif b.kind == "ul":
            for it in b.content:  # type: ignore[union-attr]
                out += ["  * " + _wrap(it, 84, indent="    ")]
            out.append("")
        elif b.kind == "ol":
            for idx, it in enumerate(b.content, 1):  # type: ignore[union-attr]
                out += [f"  {idx}. " + _wrap(it, 84, indent="     ")]
            out.append("")
        elif b.kind == "code":
            out += ["    " + line for line in str(b.content).splitlines()]
            out.append("")
        elif b.kind == "table":
            out += _render_txt_table(b.content)  # type: ignore[arg-type]
            out.append("")
    return "\n".join(out).rstrip() + "\n"


def _wrap(text: str, width: int, *, indent: str = "") -> str:
    words = text.split()
    if not words:
        return ""
    lines: list[str] = []
    current = ""
    for w in words:
        if len(current) + len(w) + 1 > width:
            lines.append(current)
            current = w
        else:
            current = f"{current} {w}".strip()
    if current:
        lines.append(current)
    return ("\n" + indent).join(lines)


def _render_txt_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    cols = max(len(r) for r in rows)
    widths = [0] * cols
    norm: list[list[str]] = []
    for r in rows:
        padded = list(r) + [""] * (cols - len(r))
        norm.append(padded)
        for i, cell in enumerate(padded):
            widths[i] = max(widths[i], len(cell))
    def fmt_row(r: list[str]) -> str:
        return "| " + " | ".join(c.ljust(widths[i]) for i, c in enumerate(r)) + " |"
    sep = "+-" + "-+-".join("-" * w for w in widths) + "-+"
    out = [sep, fmt_row(norm[0]), sep]
    for r in norm[1:]:
        out.append(fmt_row(r))
    out.append(sep)
    return out


# ----------------------------------------------------------------------- DOCX
def render_docx(blocks: list[Block], dest: Path, *, title: str, subtitle: str) -> None:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise SystemExit(
            "python-docx is required for DOCX output — `pip install python-docx`"
        ) from exc

    doc = Document()

    # Document defaults
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title_p = doc.add_paragraph()
    run = title_p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_p = doc.add_paragraph()
    sr = subtitle_p.add_run(subtitle)
    sr.italic = True
    sr.font.size = Pt(12)
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    for b in blocks:
        if b.kind == "hr":
            doc.add_paragraph("─" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if b.kind.startswith("h") and b.kind[1:].isdigit():
            level = int(b.kind[1:])
            if level == 1:
                heading = doc.add_heading(str(b.content), level=0)
            else:
                heading = doc.add_heading(str(b.content), level=min(level - 1, 4))
            for r in heading.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        elif b.kind == "p":
            doc.add_paragraph(str(b.content))
        elif b.kind == "ul":
            for it in b.content:  # type: ignore[union-attr]
                doc.add_paragraph(it, style="List Bullet")
        elif b.kind == "ol":
            for it in b.content:  # type: ignore[union-attr]
                doc.add_paragraph(it, style="List Number")
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(str(b.content))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(6)
        elif b.kind == "table":
            rows = b.content  # type: ignore[assignment]
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=cols)
            t.style = "Light Grid Accent 1"
            for r_i, row in enumerate(rows):
                cells = list(row) + [""] * (cols - len(row))
                for c_i, cell in enumerate(cells):
                    t_cell = t.rows[r_i].cells[c_i]
                    t_cell.text = ""
                    p = t_cell.paragraphs[0]
                    run = p.add_run(cell)
                    if r_i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # header shading
            try:
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement
                for cell in t.rows[0].cells:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:fill"), "1F2A44")
                    tcPr.append(shd)
            except Exception:
                pass
            doc.add_paragraph()

    doc.save(dest)


# ----------------------------------------------------------------------- PDF
def _register_unicode_fonts() -> tuple[str, str, str, str]:
    """Register a Unicode TTF family with ReportLab.

    Returns (regular, bold, italic, mono) font names. Falls back to the
    built-in Helvetica family if no system Unicode TTF can be located —
    note that Helvetica only supports Latin-1, so Cyrillic will render as
    boxes in that case.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.pdfmetrics import registerFontFamily

    candidates = [
        # (family, regular, bold, italic)
        ("DejaVuSans",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
        ("DejaVuSans",
         "C:/Windows/Fonts/DejaVuSans.ttf",
         "C:/Windows/Fonts/DejaVuSans-Bold.ttf",
         "C:/Windows/Fonts/DejaVuSans-Oblique.ttf"),
        ("Arial",
         "C:/Windows/Fonts/arial.ttf",
         "C:/Windows/Fonts/arialbd.ttf",
         "C:/Windows/Fonts/ariali.ttf"),
        ("Calibri",
         "C:/Windows/Fonts/calibri.ttf",
         "C:/Windows/Fonts/calibrib.ttf",
         "C:/Windows/Fonts/calibrii.ttf"),
        ("Tahoma",
         "C:/Windows/Fonts/tahoma.ttf",
         "C:/Windows/Fonts/tahomabd.ttf",
         None),
    ]
    mono_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
        "C:/Windows/Fonts/consola.ttf",
        "C:/Windows/Fonts/cour.ttf",
        "C:/Windows/Fonts/DejaVuSansMono.ttf",
    ]

    family = bold = italic = None
    for fam, reg, bd, it in candidates:
        if Path(reg).is_file():
            try:
                pdfmetrics.registerFont(TTFont(fam, reg))
                family = fam
                if bd and Path(bd).is_file():
                    pdfmetrics.registerFont(TTFont(f"{fam}-Bold", bd))
                    bold = f"{fam}-Bold"
                if it and Path(it).is_file():
                    pdfmetrics.registerFont(TTFont(f"{fam}-Italic", it))
                    italic = f"{fam}-Italic"
                registerFontFamily(
                    fam,
                    normal=family,
                    bold=bold or family,
                    italic=italic or family,
                    boldItalic=bold or family,
                )
                break
            except Exception:
                family = bold = italic = None
                continue

    mono = None
    for path in mono_candidates:
        if Path(path).is_file():
            try:
                pdfmetrics.registerFont(TTFont("UnicodeMono", path))
                mono = "UnicodeMono"
                break
            except Exception:
                continue

    if family is None:
        return ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier")
    return (family, bold or family, italic or family, mono or family)


def render_pdf(blocks: list[Block], dest: Path, *, title: str, subtitle: str,
               footer: str, page_word: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        ListFlowable, ListItem, PageBreak, Paragraph, Preformatted,
        SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    regular, bold, italic, mono = _register_unicode_fonts()
    global _HEADER_FONT_NAME
    _HEADER_FONT_NAME = bold

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "title", parent=styles["Title"], fontSize=22, fontName=bold,
        textColor=colors.HexColor("#1F2A44"), spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "subtitle", parent=styles["Normal"], fontSize=11, fontName=italic,
        textColor=colors.HexColor("#475569"), alignment=1, spaceAfter=18,
    )
    h_styles = {
        1: ParagraphStyle("h1", parent=styles["Heading1"], fontSize=18, fontName=bold,
                          textColor=colors.HexColor("#1F2A44"), spaceBefore=14, spaceAfter=6),
        2: ParagraphStyle("h2", parent=styles["Heading2"], fontSize=14, fontName=bold,
                          textColor=colors.HexColor("#1F2A44"), spaceBefore=10, spaceAfter=4),
        3: ParagraphStyle("h3", parent=styles["Heading3"], fontSize=12, fontName=bold,
                          textColor=colors.HexColor("#1F2A44"), spaceBefore=8,  spaceAfter=4),
        4: ParagraphStyle("h4", parent=styles["Heading4"], fontSize=11, fontName=bold,
                          textColor=colors.HexColor("#334155"), spaceBefore=6,  spaceAfter=2),
    }
    p_style = ParagraphStyle(
        "para", parent=styles["BodyText"], fontSize=10, fontName=regular,
        leading=14, spaceAfter=6,
    )
    code_style = ParagraphStyle(
        "code", parent=styles["Code"], fontSize=8.5, leading=11, fontName=mono,
        backColor=colors.HexColor("#0f172a"), textColor=colors.HexColor("#e2e8f0"),
        leftIndent=10, rightIndent=10, spaceBefore=4, spaceAfter=8,
        borderPadding=6,
    )

    doc = SimpleDocTemplate(
        str(dest), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title=title,
    )
    doc._cybersim_footer = footer
    doc._cybersim_page_word = page_word
    doc._cybersim_font = regular

    story: list = []
    story.append(Paragraph(title, title_style))
    story.append(Paragraph(subtitle, subtitle_style))
    story.append(Spacer(1, 0.4 * cm))

    for b in blocks:
        if b.kind == "hr":
            story.append(Spacer(1, 0.2 * cm))
            continue
        if b.kind.startswith("h") and b.kind[1:].isdigit():
            level = int(b.kind[1:])
            story.append(Paragraph(_pdf_escape(str(b.content)), h_styles.get(level, h_styles[4])))
        elif b.kind == "p":
            story.append(Paragraph(_pdf_escape(str(b.content)), p_style))
        elif b.kind == "ul":
            items = [ListItem(Paragraph(_pdf_escape(it), p_style)) for it in b.content]  # type: ignore[union-attr]
            story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=14))
            story.append(Spacer(1, 0.15 * cm))
        elif b.kind == "ol":
            items = [ListItem(Paragraph(_pdf_escape(it), p_style)) for it in b.content]  # type: ignore[union-attr]
            story.append(ListFlowable(items, bulletType="1", leftIndent=14))
            story.append(Spacer(1, 0.15 * cm))
        elif b.kind == "code":
            story.append(Preformatted(str(b.content), code_style))
        elif b.kind == "table":
            rows = b.content  # type: ignore[assignment]
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            normalised = [list(r) + [""] * (cols - len(r)) for r in rows]
            wrapped = [
                [Paragraph(_pdf_escape(cell), p_style if r_i else _header_para_style()) for cell in row]
                for r_i, row in enumerate(normalised)
            ]
            available_width = doc.width
            col_width = available_width / cols
            t = Table(wrapped, colWidths=[col_width] * cols, hAlign="LEFT", repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, 0), bold),
                ("FONTNAME", (0, 1), (-1, -1), regular),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94a3b8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.3 * cm))

    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)


_HEADER_FONT_NAME = "Helvetica-Bold"  # overwritten by render_pdf at runtime


def _header_para_style():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    return ParagraphStyle(
        "hdr", parent=getSampleStyleSheet()["BodyText"],
        fontSize=9.5, leading=12, textColor=colors.whitesmoke,
        fontName=_HEADER_FONT_NAME,
    )


def _pdf_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
    )


def _pdf_footer(canvas, doc):
    from reportlab.lib import colors
    footer = getattr(doc, "_cybersim_footer", "CyberSim Documentation — v0.1.0")
    page_word = getattr(doc, "_cybersim_page_word", "Page")
    font_name = getattr(doc, "_cybersim_font", "Helvetica")
    canvas.saveState()
    try:
        canvas.setFont(font_name, 8)
    except Exception:
        canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#94a3b8"))
    canvas.drawString(2 * 28.35, 1 * 28.35, footer)
    canvas.drawRightString(
        canvas._pagesize[0] - 2 * 28.35, 1 * 28.35,
        f"{page_word} {doc.page}",
    )
    canvas.restoreState()


# ----------------------------------------------------------------------- main
def _build_one(*, source: Path, prefix: str, title: str, subtitle: str,
               footer: str, page_word: str) -> int:
    if not source.is_file():
        print(f"  source missing: {source}", file=sys.stderr)
        return 1
    md = source.read_text(encoding="utf-8")
    blocks = parse_markdown(md)

    out_txt  = source.parent / f"{prefix}.txt"
    out_docx = source.parent / f"{prefix}.docx"
    out_pdf  = source.parent / f"{prefix}.pdf"

    print(f"[{prefix}] parsed {len(blocks)} blocks from {source.name}")

    out_txt.write_text(render_txt(blocks), encoding="utf-8")
    print(f"  wrote {out_txt.name} ({out_txt.stat().st_size:,} bytes)")

    try:
        render_docx(blocks, out_docx, title=title, subtitle=subtitle)
        print(f"  wrote {out_docx.name} ({out_docx.stat().st_size:,} bytes)")
    except SystemExit as exc:
        print(f"  DOCX skipped: {exc}")

    try:
        render_pdf(blocks, out_pdf, title=title, subtitle=subtitle,
                   footer=footer, page_word=page_word)
        print(f"  wrote {out_pdf.name} ({out_pdf.stat().st_size:,} bytes)")
    except Exception as exc:  # noqa: BLE001
        print(f"  PDF failed: {exc}")
        raise

    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build CyberSim docs in TXT/DOCX/PDF.")
    parser.add_argument(
        "--lang", action="append", choices=sorted(LANGS.keys()),
        help="Language preset to build (en, ru). Repeat for multiple. "
             "Default: en if no --source is given.",
    )
    parser.add_argument("--source", type=Path, help="Custom markdown source path.")
    parser.add_argument("--prefix", help="Custom output filename prefix.")
    parser.add_argument("--title", help="Custom document title.")
    parser.add_argument("--subtitle", help="Custom document subtitle.")
    parser.add_argument("--footer", help="Custom PDF footer text.")
    parser.add_argument("--page-word", default="Page",
                        help="Localised word for 'Page' in the PDF footer.")
    args = parser.parse_args(argv)

    if args.source:
        prefix = args.prefix or args.source.stem
        return _build_one(
            source=args.source.resolve(),
            prefix=prefix,
            title=args.title or prefix,
            subtitle=args.subtitle or "",
            footer=args.footer or prefix,
            page_word=args.page_word,
        )

    langs = args.lang or ["en"]
    rc = 0
    for lang in langs:
        preset = LANGS[lang]
        rc |= _build_one(
            source=preset["source"],
            prefix=preset["prefix"],
            title=preset["title"],
            subtitle=preset["subtitle"],
            footer=preset["footer"],
            page_word=preset["page_word"],
        )
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
